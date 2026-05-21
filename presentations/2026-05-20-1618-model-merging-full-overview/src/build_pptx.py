#!/usr/bin/env python3
"""Build a native PowerPoint version of the model-merging full overview."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "figures"
OUT = ROOT / "model_merging_full_overview.pptx"

TITLE = RGBColor(31, 45, 61)
BLUE = RGBColor(48, 102, 190)
GREEN = RGBColor(0, 166, 118)
ORANGE = RGBColor(224, 122, 95)
GRAY = RGBColor(82, 95, 110)
LIGHT = RGBColor(246, 248, 251)


def add_textbox(slide, x, y, w, h, text, *, size=24, bold=False, color=TITLE, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = text
    if align:
        p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def title(slide, text):
    add_textbox(slide, 0.55, 0.28, 12.2, 0.55, text, size=24, bold=True)
    line = slide.shapes.add_shape(1, Inches(0.55), Inches(0.88), Inches(12.2), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.color.rgb = BLUE


def bullet_slide(prs, heading, bullets, *, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, heading)
    shape = slide.shapes.add_textbox(Inches(0.85), Inches(1.2), Inches(11.75), Inches(5.55))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.clear()
    for i, item in enumerate(bullets):
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = text
        p.level = level
        p.font.size = Pt(22 if level == 0 else 18)
        p.font.color.rgb = TITLE if level == 0 else GRAY
        p.space_after = Pt(7)
    if note:
        add_textbox(slide, 0.85, 6.8, 11.7, 0.35, note, size=13, color=GRAY)
    return slide


def image_slide(prs, heading, image_name, *, caption=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, heading)
    slide.shapes.add_picture(str(FIG / image_name), Inches(0.85), Inches(1.18), width=Inches(11.6))
    if caption:
        add_textbox(slide, 0.85, 6.85, 11.6, 0.35, caption, size=13, color=GRAY, align=PP_ALIGN.CENTER)
    return slide


def table_slide(prs, heading, headers, rows, *, font_size=12):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, heading)
    table = slide.shapes.add_table(
        len(rows) + 1,
        len(headers),
        Inches(0.55),
        Inches(1.15),
        Inches(12.25),
        Inches(5.75),
    ).table
    widths = [1.6] + [10.65 / (len(headers) - 1)] * (len(headers) - 1) if len(headers) > 1 else [12.25]
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.bold = True
                r.font.size = Pt(font_size)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT if i % 2 == 0 else RGBColor(255, 255, 255)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font_size)
                    r.font.color.rgb = TITLE
    return slide


def section_slide(prs, heading, subheading):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(245, 248, 252)
    add_textbox(slide, 0.85, 2.25, 11.7, 0.75, heading, size=36, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    add_textbox(slide, 1.45, 3.15, 10.5, 0.65, subheading, size=20, color=GRAY, align=PP_ALIGN.CENTER)
    return slide


def build() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, 0.75, 1.95, 11.9, 0.75, "Mechanistic Model Merging", size=40, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    add_textbox(slide, 1.05, 2.85, 11.25, 0.55, "Plain-language overview, project plan, evidence so far, and current checkpoint", size=20, color=GRAY, align=PP_ALIGN.CENTER)
    add_textbox(slide, 1.05, 4.2, 11.25, 0.45, "Gavin + Codex | May 20, 2026", size=16, color=TITLE, align=PP_ALIGN.CENTER)

    bullet_slide(prs, "One-Slide Summary", [
        "Model merging combines trained model weights or fine-tuning deltas into one model.",
        "Our project asks why merging works or fails mechanistically, not only whether benchmark scores move.",
        "Current case: public Qwen2.5-1.5B abliterated TIES merge loses harmful-refusal behavior.",
        "Main checkpoint: low-rank MLP activation-delta subspace over layers 12-24 repairs the behavior.",
        "Current baseline to beat: PCA rank 64, after strict manual audit.",
    ])

    section_slide(prs, "Part 1: What Model Merging Is", "Simple version first, then why it is useful and hard.")
    bullet_slide(prs, "What Is Model Merging?", [
        "A neural network is a large table of learned numbers: weights.",
        "Merging combines checkpoints or fine-tuning deltas into one new checkpoint.",
        "It happens before inference; it is not model voting or routing at runtime.",
        "The usual practical hope: inherit useful behaviors without full retraining.",
    ])
    bullet_slide(prs, "The Simplest Formula", [
        "Expert = base + task delta.",
        "Merged = base + alpha * math_delta + beta * code_delta + gamma * safety_delta.",
        "This is weight-space editing: we change the model itself.",
        "It works best when models share a base or have aligned internal coordinates.",
    ])
    image_slide(prs, "Merging Is Not Concatenation", "merge_concept.png")
    bullet_slide(prs, "Why People Use It", [
        "Only checkpoints are available; original training data may be missing, private, or licensed.",
        "Full joint training is too expensive.",
        "One merged model is cheaper to serve than many experts.",
        "Open-source LLM communities use it for math, code, chat, multilingual, and safety variants.",
        "The realistic claim is often: best possible with checkpoints only, not better than ideal full-data training.",
    ])
    table_slide(prs, "Related Work Anchors", ["Work", "Concept", "Why it matters"], [
        ["FedAvg, 2017", "average client updates", "checkpoint/update averaging ancestor"],
        ["SWA, 2018", "average training checkpoints", "averaging can find wider optima"],
        ["Model soups, 2022", "average fine-tuned models", "modern same-base merging"],
        ["Task Arithmetic, 2023", "add/subtract task vectors", "delta view of editing"],
        ["TIES, 2023", "trim and resolve signs", "interference handling"],
        ["DARE, 2024", "drop/rescale deltas", "delta redundancy"],
        ["MergeKit, 2024", "LLM merge toolkit", "practical open-source workflows"],
    ], font_size=11)
    bullet_slide(prs, "Why It Is Hard To Interpret", [
        "Good source models may use incompatible internal coordinates.",
        "Deltas can interfere by sign, magnitude, or layer.",
        "A behavior may require detection, policy, wording, coherence, and termination.",
        "A merge can preserve shallow keywords but lose the full mechanism.",
        "This is why mechanistic tests and manual audits matter.",
    ])

    section_slide(prs, "Part 2: What This Project Wants To Do", "Turn merging from score comparison into mechanistic explanation.")
    bullet_slide(prs, "Central Project Question", [
        "When model merging works or fails, can we explain it in terms of identifiable mechanisms?",
        "Mechanisms can mean modules, activation directions, sparse features, or causal pathways.",
        "The goal is not just to make a better merge.",
        "The goal is to understand what was inherited, what was lost, and why.",
    ])
    image_slide(prs, "Project Pipeline", "project_pipeline.png")
    table_slide(prs, "Research Questions RQ0-RQ7", ["RQ", "Question"], [
        ["RQ0", "When do simple baselines already explain merging?"],
        ["RQ1", "What predicts merge success before sparse features?"],
        ["RQ2", "Are inherited capabilities carried by the same modules?"],
        ["RQ3", "Which modules are necessary and sufficient?"],
        ["RQ4", "Why are some restored behaviors messy?"],
        ["RQ5", "What is merge interference mechanistically?"],
        ["RQ6", "Do expert deltas correspond to mechanisms?"],
        ["RQ7", "Are SAE/transcoder features a better basis?"],
    ], font_size=11)
    table_slide(prs, "Research Questions RQ8-RQ15", ["RQ", "Question"], [
        ["RQ8", "What feature types exist during merging?"],
        ["RQ9", "Can sparse feature patching recover missing capabilities?"],
        ["RQ10", "Can mechanism-aware merging outperform naive merging?"],
        ["RQ11", "Does the explanation generalize?"],
        ["RQ12", "How do merge algorithms differ mechanistically?"],
        ["RQ13", "Can failures be predicted before evaluation?"],
        ["RQ14", "Does merging create new composition?"],
        ["RQ15", "What should practitioners do differently?"],
    ], font_size=11)
    table_slide(prs, "Experimental Phases", ["Phase", "Goal", "Status"], [
        ["A", "strengthen behavior benchmark", "SmolLM2 taught evaluator caution"],
        ["B", "basis-validation benchmark", "active on Qwen2.5-1.5B"],
        ["C", "SAE/transcoder training/loading", "gated"],
        ["D", "sparse feature discovery", "future"],
        ["E", "causal feature validation", "future"],
        ["F", "mechanism-aware merge recipes", "future"],
        ["G", "generalization and paper claims", "future"],
    ], font_size=11)

    section_slide(prs, "Part 3: What We Have Done", "From toy merges to a public Qwen safety-loss case.")
    bullet_slide(prs, "Stage 0: Controlled Sandbox", [
        "Built MNIST domain experts and SmolLM2-135M synthetic experts.",
        "Tested arithmetic, politeness, and refusal merges.",
        "Confirmed simple merging can compose toy behaviors.",
        "But toy behavior was not enough for a strong LLM mechanism paper.",
    ])
    bullet_slide(prs, "SmolLM2 Branch: Useful Failure", [
        "Refusal-like fragments transferred.",
        "Clean refusal was rare.",
        "Outputs often had repetition, artifacts, contradiction, or unsafe continuation.",
        "Decision: do not spend expensive SAE work there as the main clean-safety case.",
    ])
    bullet_slide(prs, "Public Merge Screening", [
        "Screened Qwen-family public merges.",
        "Many 0.5B candidates were near-copies or behaviorally weak.",
        "Qwen2.5-1.5B gave a stronger public case.",
        "Current pair: Qwen2.5-1.5B-Instruct vs EVA-abliterated-TIES-Qwen2.5-1.5B.",
    ])
    bullet_slide(prs, "Why The Qwen Pair Is Good", [
        "Base refuses harmful prompts.",
        "Abliterated model gives harmful instructions.",
        "Benign helpfulness is mostly preserved.",
        "Prompt-specific activation drift is visible.",
        "This creates a merge-induced safety-loss case study.",
    ])

    section_slide(prs, "Part 4: Mechanistic Evidence So Far", "Localization, causal repair, then compressed baselines.")
    image_slide(prs, "RQ1/RQ2: Activation Drift", "activation_drift.png", caption="Harmful prompts drift more than benign prompts in mid-to-late layers.")
    bullet_slide(prs, "Single Modules Were Not Enough", [
        "Target-loss patching pointed to mid-layer MLP/block modules.",
        "But generation validation failed for single modules.",
        "This separated likelihood localization from behavioral sufficiency.",
        "Conclusion: refusal repair is distributed.",
    ])
    image_slide(prs, "RQ3: Broad MLP Range Repairs Behavior", "localization_range.png", caption="Static MLP replacement needs a broad contiguous range; layers 12-24 are strongest.")
    bullet_slide(prs, "Activation Patch Confirms Mechanism Level", [
        "12-24 MLP activation patch closes 0.938 of harmful refusal-target loss gap.",
        "12-21 closes 0.937; 12-19 closes 0.917.",
        "Target-position-only patching closes 0.000.",
        "Repair requires sequence-wide MLP activation propagation.",
    ])
    bullet_slide(prs, "Dynamic Generation Patch", [
        "Run donor/base on current prefix.",
        "Patch donor MLP activations into abliterated recipient during greedy decoding.",
        "12-24 MLP patch restores harmful refusal to 1.000 on 8 harmful prompts.",
        "Benign helpfulness remains 1.000; over-refusal remains 0.000.",
    ])

    section_slide(prs, "Part 5: Today's Main Result", "PCA rank 64 is the current compressed repair baseline.")
    bullet_slide(prs, "RQ0: Simple Bases Before SAE", [
        "Tested mean donor-recipient delta direction per layer.",
        "Tested top changed neurons.",
        "Tested harmful-vs-benign refusal direction.",
        "Tested randomized PCA over activation deltas.",
        "Tested matched random projection controls.",
    ])
    image_slide(prs, "Target-Loss Repair: PCA Wins", "pca_target_loss.png", caption="PCA rank 64 closes 0.955 of the harmful target-loss gap; random rank 64 closes 0.230.")
    image_slide(prs, "Behavior: Strict Audit Matters", "pca_generation_audit.png", caption="Automatic scoring over-credits some refusal-prefix answers; strict manual audit keeps PCA rank 64 strongest.")
    table_slide(prs, "Manual Audit Result", ["Patch", "Auto harmful", "Strict harmful", "Benign helpful"], [
        ["base", "0.875", "1.000", "1.000"],
        ["abliterated", "0.000", "0.000", "1.000"],
        ["mean_delta_rank1", "0.875", "0.750", "1.000"],
        ["pca_rank16", "1.000", "0.625", "1.000"],
        ["pca_rank64", "1.000", "1.000", "1.000"],
        ["random_rank64", "0.000", "0.000", "1.000"],
    ], font_size=12)
    bullet_slide(prs, "Current Interpretation", [
        "The abliterated TIES merge appears to lose a low-dimensional MLP refusal-restoration subspace.",
        "The relevant pathway spans layers 12-24 MLP outputs.",
        "Mean-delta rank 1 is a partial minimal repair.",
        "PCA rank 64 is the strongest compressed behavioral repair.",
        "SAE/transcoder features must explain, compress, or improve on this PCA subspace.",
    ])
    bullet_slide(prs, "What We Should Not Claim Yet", [
        "Not a proof that merging is generally safe or unsafe.",
        "Not a full safety benchmark.",
        "Not an SAE result yet.",
        "Not a solved mechanism; PCA rank 64 is a strong baseline, not an explanation by itself.",
    ])

    section_slide(prs, "Part 6: Next Steps", "Harden the target, then test sparse features.")
    bullet_slide(prs, "Immediate Next Steps", [
        "Fix the harmful-refusal evaluator to detect unsafe continuation after refusal prefixes.",
        "Build held-out harmful and adversarial benign prompt sets.",
        "Revalidate PCA rank 64, mean-delta, full MLP patch, and random controls.",
        "Compare exact/larger-sample PCA to current randomized PCA.",
        "Only then start SAE/transcoder feature experiments.",
    ])
    bullet_slide(prs, "What SAE Must Beat", [
        "Better repair than PCA rank 64 at similar or lower dimensionality.",
        "Interpretable decomposition of the PCA repair into feature groups.",
        "Causal sparse feature patches that recover refusal without unsafe continuation.",
        "Prediction of which merges will preserve or lose refusal before generation tests.",
    ])
    bullet_slide(prs, "Paper Potential", [
        "Current status: continue, do not abandon.",
        "Strengths: public model case, causal repair, strong baseline, clear evaluator caveat.",
        "Main risk: sparse features may not beat PCA or add insight.",
        "If that happens, frame as low-rank activation geometry rather than SAE paper.",
    ])
    bullet_slide(prs, "Artifact Map", [
        "Package: presentations/2026-05-20-1618-model-merging-full-overview/",
        "PDF: model_merging_full_overview.pdf",
        "PowerPoint: model_merging_full_overview.pptx",
        "Data: copied CSV metrics and manual audit in data/",
        "Figures: generated from copied data in figures/",
    ])
    bullet_slide(prs, "Bottom Line", [
        "We now have a credible model-merging interpretability target.",
        "It is a merge-induced refusal failure.",
        "It can be repaired by restoring a low-rank MLP activation-delta subspace.",
        "The next decisive question is whether sparse features explain this better than PCA rank 64.",
    ])

    prs.save(OUT)


if __name__ == "__main__":
    build()
