#!/usr/bin/env python3
from __future__ import annotations

import csv
from pathlib import Path

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
FIGURES = ROOT / "figures"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def row_by_model(rows: list[dict[str, str]], model: str) -> dict[str, str]:
    for row in rows:
        if row["model"] == model:
            return row
    raise KeyError(model)


def metric(row: dict[str, str], key: str) -> float:
    return float(row[key])


def build_combined_rows() -> list[dict[str, object]]:
    post_ff = read_csv(DATA / "post_ff_activation_patch_metrics.csv")
    sae_16_20 = read_csv(DATA / "mlp_sae_16_20_generation_metrics.csv")
    sae_12_20 = read_csv(DATA / "mlp_sae_12_20_generation_metrics.csv")
    specs = [
        ("Base", row_by_model(post_ff, "base")),
        ("Abliterated", row_by_model(post_ff, "abliterated")),
        ("Full post-FF 16-20", row_by_model(post_ff, "activation_patch_16+17+18+19+20:post_ff")),
        ("Full post-FF 12-20", row_by_model(post_ff, "activation_patch_12+13+14+15+16+17+18+19+20:post_ff")),
        ("SAE decoded 16-20", row_by_model(sae_16_20, "gemmascope_mlp_sae_patch_16+17+18+19+20")),
        ("SAE decoded 12-20", row_by_model(sae_12_20, "gemmascope_mlp_sae_patch_12+13+14+15+16+17+18+19+20")),
    ]
    rows = []
    for label, row in specs:
        rows.append(
            {
                "label": label,
                "harmful_clean": metric(row, "harmful_ok_rate"),
                "unsafe": metric(row, "harmful_unsafe_continuation_rate"),
                "benign_helpful": metric(row, "benign_ok_rate"),
            }
        )
    return rows


def write_combined_csv(rows: list[dict[str, object]]) -> Path:
    path = DATA / "combined_behavior_gate.csv"
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["label", "harmful_clean", "unsafe", "benign_helpful"])
        writer.writeheader()
        writer.writerows(rows)
    return path


def build_behavior_plot(rows: list[dict[str, object]]) -> Path:
    FIGURES.mkdir(parents=True, exist_ok=True)
    labels = [str(row["label"]) for row in rows]
    x = range(len(rows))
    width = 0.24
    fig, ax = plt.subplots(figsize=(11, 5.8))
    ax.bar([i - width for i in x], [float(row["harmful_clean"]) for row in rows], width, label="harmful clean", color="#2f6f73")
    ax.bar(list(x), [float(row["benign_helpful"]) for row in rows], width, label="benign helpful", color="#6b8e23")
    ax.bar([i + width for i in x], [float(row["unsafe"]) for row in rows], width, label="unsafe continuation", color="#b44a4a")
    ax.set_ylim(0, 1.05)
    ax.set_ylabel("Rate on 4 harmful + 4 benign prompt screen")
    ax.set_title("GemmaScope MLP SAE decoded patch matches the full 12-20 repair gate")
    ax.set_xticks(list(x))
    ax.set_xticklabels(labels, rotation=25, ha="right")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(loc="lower right")
    fig.tight_layout()
    path = FIGURES / "behavior_gate.png"
    fig.savefig(path, dpi=180)
    return path


def set_title(slide, title: str):
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(30)


def add_bullets(slide, bullets: list[str], left=0.7, top=1.35, width=11.9, height=5.0):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.level = 0
        p.font.size = Pt(20)


def add_table(slide, headers: list[str], rows: list[list[str]], left, top, width, height):
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(left), Inches(top), Inches(width), Inches(height))
    table = shape.table
    for c, header in enumerate(headers):
        table.cell(0, c).text = header
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    for row in table.rows:
        for cell in row.cells:
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(12)
    return shape


def build_pptx(plot_path: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "GemmaScope MLP SAE Checkpoint"
    slide.placeholders[1].text = "Model merging SAE project - 2026-05-22"

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(slide, "Question")
    add_bullets(
        slide,
        [
            "Can a public sparse basis reproduce a causal model-merging repair, not just describe activations?",
            "Target pair: Gemma-2-2B-IT base donor into an abliterated recipient.",
            "Behavior gate: restore harmful refusal while preserving benign helpfulness.",
            "Important distinction: this is decoded SAE completeness, not yet feature-level explanation.",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(slide, "Causal Target")
    add_table(
        slide,
        ["Patch", "Harmful clean", "Unsafe", "Benign helpful"],
        [
            ["Abliterated", "0.000", "0.000", "1.000"],
            ["Full post-FF 16-20", "0.750", "0.000", "1.000"],
            ["Full post-FF 12-20", "1.000", "0.000", "1.000"],
        ],
        0.8,
        1.4,
        11.6,
        1.7,
    )
    add_bullets(
        slide,
        [
            "The successful activation site is the post-feedforward MLP update, layers 12-20.",
            "This aligns the causal target with GemmaScope MLP SAEs.",
        ],
        top=3.5,
        height=2.5,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(slide, "Sparse-Basis Gate")
    slide.shapes.add_picture(str(plot_path), Inches(0.7), Inches(1.15), width=Inches(12.0))

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(slide, "Interpretation")
    add_bullets(
        slide,
        [
            "Positive: GemmaScope MLP-SAE decoded 12-20 patch passes the same behavior gate as full activation patching.",
            "Negative control: the corrected layer-16 transcoder is weaker, so MLP SAE is the priority branch.",
            "Caveat: top_neuron_k1536 also passes on this small screen, so feature selection is still required for novelty.",
            "Scientific move: go from full decoded reconstruction to sparse feature subsets with matched random controls.",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(slide, "Next Step")
    add_bullets(
        slide,
        [
            "Run donor-active and donor-recipient feature-delta patching inside layers 12-20.",
            "Compare feature subsets against top-coordinate and random active-feature baselines.",
            "Expand to held-out harmful prompts and adversarial benign controls before claiming mechanism.",
            "If subsets fail, the honest claim is behavioral completeness without sparse feature-level explanation.",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    set_title(slide, "Provenance")
    add_bullets(
        slide,
        [
            "stage2/results/gemma2_2b_dynamic_activation_patch_generation_post_ff/",
            "stage3/results/gemma2_2b_gemmascope_mlp_sae_validation_12_20_generation/",
            "stage3/scripts/validate_gemma2_2b_gemmascope_mlp_sae.py",
            "stage3/results/PROVENANCE.md",
        ],
    )

    out = ROOT / "gemmascope_mlp_sae_checkpoint.pptx"
    prs.save(out)
    return out


def main():
    rows = build_combined_rows()
    write_combined_csv(rows)
    plot_path = build_behavior_plot(rows)
    build_pptx(plot_path)


if __name__ == "__main__":
    main()
