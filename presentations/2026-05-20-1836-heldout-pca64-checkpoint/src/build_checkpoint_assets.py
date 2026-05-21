#!/usr/bin/env python3
"""Build assets for the held-out PCA64 checkpoint presentation."""

from __future__ import annotations

import csv
from pathlib import Path

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
FIG = ROOT / "figures"
FIG.mkdir(parents=True, exist_ok=True)


def read_metrics(path: Path) -> dict[str, dict[str, float]]:
    out = {}
    with path.open() as f:
        for row in csv.DictReader(f):
            out[row["model"]] = {k: float(v) if k != "model" else v for k, v in row.items()}
    return out


def build_plot() -> None:
    heldout = read_metrics(DATA / "pca64_heldout_metrics.csv")
    calib12 = read_metrics(DATA / "pca64_heldout_calib12_metrics.csv")
    full = read_metrics(DATA / "full_patch_heldout_metrics.csv")
    rows = [
        ("base", heldout["base"]),
        ("abliterated", heldout["abliterated"]),
        ("PCA64", heldout["pca_rank64"]),
        ("PCA64 calib12", calib12["pca_rank64"]),
        ("random64", heldout["random_rank64"]),
        ("full MLP patch", full["full_12_24_mlp"]),
    ]
    labels = [name for name, _row in rows]
    harmful = [row["harmful_ok_rate"] for _name, row in rows]
    benign = [row["benign_ok_rate"] for _name, row in rows]
    unsafe = [row["harmful_unsafe_continuation_rate"] for _name, row in rows]

    fig, ax = plt.subplots(figsize=(9.2, 3.8))
    x = list(range(len(rows)))
    width = 0.27
    ax.bar([i - width for i in x], harmful, width, label="strict harmful clean", color="#3066BE")
    ax.bar(x, benign, width, label="benign helpful", color="#00A676")
    ax.bar([i + width for i in x], unsafe, width, label="unsafe content", color="#E07A5F")
    ax.set_xticks(x, labels, rotation=15, ha="right")
    ax.set_ylim(0, 1.08)
    ax.set_ylabel("rate")
    ax.set_title("Held-out validation: full patch remains clean; PCA64 is partial")
    ax.legend(frameon=False, ncol=3, loc="lower center", bbox_to_anchor=(0.5, -0.34))
    plt.tight_layout()
    plt.savefig(FIG / "heldout_patch_comparison.pdf")
    plt.savefig(FIG / "heldout_patch_comparison.png", dpi=180)
    plt.close()


def textbox(slide, x, y, w, h, text, *, size=24, bold=False, color=None, align=None):
    color = color or RGBColor(31, 45, 61)
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    if align:
        p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def title(slide, text):
    textbox(slide, 0.55, 0.28, 12.2, 0.55, text, size=24, bold=True)


def bullets(prs, heading, items):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, heading)
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.8), Inches(5.8))
    tf = box.text_frame
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(22)
        p.space_after = Pt(8)
    return slide


def build_pptx() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    textbox(slide, 0.9, 2.05, 11.5, 0.8, "Held-out PCA64 Checkpoint", size=38, bold=True, color=RGBColor(48, 102, 190), align=PP_ALIGN.CENTER)
    textbox(slide, 1.0, 3.0, 11.3, 0.5, "Model merging refusal repair: strict held-out validation", size=20, align=PP_ALIGN.CENTER)
    textbox(slide, 1.0, 4.1, 11.3, 0.35, "May 20, 2026", size=16, align=PP_ALIGN.CENTER)

    bullets(prs, "Why This Checkpoint Matters", [
        "The previous result made PCA rank 64 look like the strongest compressed repair.",
        "We needed held-out harmful prompts and adversarial benign prompts before trusting it.",
        "The stricter scorer rejects refusal-prefix answers that continue into unsafe advice.",
    ])
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "Held-out Result")
    slide.shapes.add_picture(str(FIG / "heldout_patch_comparison.png"), Inches(0.7), Inches(1.05), width=Inches(12.0))
    bullets(prs, "Interpretation", [
        "PCA64 remains a strong non-sparse baseline: 10/12 strict harmful refusals, benign helpfulness 12/12.",
        "But full 12-24 MLP activation patch gets 12/12 strict harmful refusals.",
        "More calibration examples did not help PCA64; it dropped to 9/12.",
        "Therefore PCA64 is not the full repair mechanism.",
    ])
    bullets(prs, "Next Step", [
        "Study the residual between full MLP activation patch and PCA64 patch.",
        "Localize which prompts/layers/directions account for PCA64 failures.",
        "Use full patch as the behavioral upper bound and PCA64 as the strongest compressed baseline.",
        "SAE/transcoder work should target the residual mechanism, not just reproduce PCA64.",
    ])
    prs.save(ROOT / "heldout_pca64_checkpoint.pptx")


def main() -> int:
    build_plot()
    build_pptx()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
