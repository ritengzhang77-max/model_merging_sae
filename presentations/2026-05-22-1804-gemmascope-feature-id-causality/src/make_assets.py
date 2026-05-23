#!/usr/bin/env python3
"""Build figures and PowerPoint for the GemmaScope feature-ID checkpoint."""

from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
FIG_DIR = ROOT / "figures"


def read_metrics() -> list[dict[str, str]]:
    with (DATA_DIR / "feature_id_threshold_metrics.csv").open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def metric(rows: list[dict[str, str]], family: str, condition: str) -> tuple[float, float]:
    for row in rows:
        if row["family"] == family and row["condition"] == condition:
            return float(row["harmful_clean"]), float(row["unsafe"])
    raise KeyError((family, condition))


def make_main_plot(rows: list[dict[str, str]]) -> Path:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    path = FIG_DIR / "l19_tail_causality.png"
    conditions = [
        ("k896", "prefix_8_12_hi", "mix_decode_delta_abs_k896"),
        ("k1024", "k1024_8_12", "mix_decode_delta_abs_k1024"),
        ("tail only\n897-1024", "tail_only", "mix_decode_delta_abs_rank897_1024"),
        ("k896 +\nL19 tail", "plus_l16_20", "mix_decode_delta_abs_k896_plus_l19_rank897_1024"),
        ("k1024 -\nL19 tail", "minus_l16_20", "mix_decode_delta_abs_k1024_minus_l19_rank897_1024"),
    ]
    vals = [metric(rows, family, condition)[0] for _, family, condition in conditions]
    colors = ["#4a5568", "#2b6cb0", "#718096", "#2f855a", "#c05621"]
    fig, ax = plt.subplots(figsize=(8.6, 4.3))
    bars = ax.bar(range(len(conditions)), vals, color=colors)
    for bar, val in zip(bars, vals, strict=True):
        ax.text(bar.get_x() + bar.get_width() / 2, val + 0.03, f"{val:.2f}", ha="center", fontsize=10)
    ax.set_ylim(0, 1.05)
    ax.set_ylabel("Harmful clean refusal rate")
    ax.set_xticks(range(len(conditions)))
    ax.set_xticklabels([label for label, _, _ in conditions])
    ax.set_title("Layer-19 rank897-1024 tail explains the k896 to k1024 fake-ID gain")
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)
    return path


def make_token_plot() -> Path:
    path = FIG_DIR / "l19_tail_event_tokens.png"
    counts: Counter[str] = Counter()
    with (DATA_DIR / "l19_tail_feature_events.jsonl").open(encoding="utf-8") as handle:
        for line in handle:
            row = json.loads(line)
            counts[str(row["token"])] += 1
    top = counts.most_common(8)
    labels = [token if token != "\\n" else "newline" for token, _ in top]
    vals = [count for _, count in top]
    fig, ax = plt.subplots(figsize=(8.6, 3.8))
    bars = ax.bar(range(len(vals)), vals, color="#805ad5")
    for bar, val in zip(bars, vals, strict=True):
        ax.text(bar.get_x() + bar.get_width() / 2, val + 5, str(val), ha="center", fontsize=9)
    ax.set_xticks(range(len(labels)))
    ax.set_xticklabels(labels, rotation=20, ha="right")
    ax.set_ylabel("Audit event count")
    ax.set_title("L19 tail feature events remain boundary/template dominated")
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)
    return path


def add_title(slide, title: str, subtitle: str = "") -> None:
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.2), Inches(0.55))
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(27)
    p.font.bold = True
    if subtitle:
        sbox = slide.shapes.add_textbox(Inches(0.48), Inches(0.88), Inches(12.2), Inches(0.35))
        sframe = sbox.text_frame
        sframe.clear()
        p = sframe.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(13)


def add_bullets(slide, bullets: list[str], top: float = 1.35, size: int = 20) -> None:
    box = slide.shapes.add_textbox(Inches(0.75), Inches(top), Inches(11.9), Inches(5.5))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    for idx, item in enumerate(bullets):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.space_after = Pt(8)


def add_table(slide, rows: list[list[str]], left: float, top: float, width: float, height: float) -> None:
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width), Inches(height))
    table = shape.table
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            for para in cell.text_frame.paragraphs:
                para.alignment = PP_ALIGN.CENTER
                for run in para.runs:
                    run.font.size = Pt(10 if r_idx else 11)
                    run.font.bold = r_idx == 0


def make_pptx(main_plot: Path, token_plot: Path) -> Path:
    pptx_path = ROOT / "gemmascope_feature_id_causality_checkpoint.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    add_title(slide, "GemmaScope Feature-ID Causality", "Model merging SAE project - 2026-05-22")
    add_bullets(
        slide,
        [
            "Checkpoint reason: position masks found the path; this run localizes one feature band inside it.",
            "Causal path: assistant/template boundary plus generated-token history.",
            "Main lead: layer 19 rank897-1024 features explain the k896 to k1024 fake-ID recovery.",
        ],
    )

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Evidence Table")
    add_table(
        slide,
        [
            ["Condition on 8:12", "Clean", "Unsafe", "Fake-ID"],
            ["k896", "0.50", "0.00", "fail"],
            ["k1024", "0.75", "0.00", "pass"],
            ["tail 897-1024 only", "0.00", "0.00", "fail"],
            ["k896 + L19 tail", "0.75", "0.25", "pass"],
            ["k1024 - L19 tail", "0.50", "0.00", "fail"],
            ["k1024 - other single tail", "0.75", "0.00", "pass"],
        ],
        0.75,
        1.35,
        11.8,
        3.2,
    )
    add_bullets(slide, ["Clean = harmful clean refusal rate; unsafe = unsafe continuation rate."], top=5.1, size=17)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Causal Pattern")
    slide.shapes.add_picture(str(main_plot), Inches(1.0), Inches(1.25), width=Inches(11.2))

    slide = prs.slides.add_slide(blank)
    add_title(slide, "What Are The L19 Tail Features?")
    slide.shapes.add_picture(str(token_plot), Inches(1.0), Inches(1.25), width=Inches(11.0))

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Interpretation")
    add_bullets(
        slide,
        [
            "Not a tiny standalone refusal module: tail-only bands score 0.00.",
            "The L19 tail is add-on sufficient and necessary for one hard prompt when the k896 prefix is present.",
            "Events remain boundary/template dominated, so the likely object is response-state refinement rather than harmful-content semantics.",
            "Next test: split L19 ranks 897-1024 into smaller blocks and inspect exact feature rows.",
        ],
    )

    prs.save(pptx_path)
    return pptx_path


def main() -> int:
    rows = read_metrics()
    main_plot = make_main_plot(rows)
    token_plot = make_token_plot()
    pptx_path = make_pptx(main_plot, token_plot)
    print(main_plot)
    print(token_plot)
    print(pptx_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
