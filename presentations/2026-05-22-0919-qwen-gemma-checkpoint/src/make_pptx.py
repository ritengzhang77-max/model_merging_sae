#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]


def add_title(slide, title: str, subtitle: str = ""):
    slide.shapes.title.text = title
    if subtitle:
        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(0.6))
        box.text_frame.text = subtitle
        box.text_frame.paragraphs[0].font.size = Pt(20)


def add_bullets(slide, bullets: list[str], top=2.1):
    box = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11.5), Inches(4.6))
    tf = box.text_frame
    tf.clear()
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(20)


def add_picture_slide(prs, title: str, figure: str, note: str):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    slide.shapes.add_picture(str(ROOT / "figures" / figure), Inches(1.2), Inches(1.25), width=Inches(11.0))
    box = slide.shapes.add_textbox(Inches(1.1), Inches(6.3), Inches(11.2), Inches(0.5))
    box.text_frame.text = note
    box.text_frame.paragraphs[0].font.size = Pt(14)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Model Merging SAE Checkpoint"
    slide.placeholders[1].text = "Qwen generated-trace SAE result and Gemma candidate promotion\n2026-05-22"

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Current Question")
    add_bullets(
        slide,
        [
            "Gate: can sparse/transcoder bases explain merge safety loss better than raw/PCA/top-coordinate/module baselines?",
            "Qwen: generated-token SAE training changes the family repaired, but does not pass the full v0 gate.",
            "Gemma: clean abliterated pair with harmful-specific activation drift and public GemmaScope bases.",
        ],
    )

    add_picture_slide(
        prs,
        "Qwen: Distribution Matters, But Gate Still Fails",
        "qwen_generated_trace_harmful_clean.png",
        "Generated-token d512 repairs one-time-code and permission-slip; tracking still fails.",
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Qwen Interpretation")
    add_bullets(
        slide,
        [
            "High SAE EV is not causal completeness.",
            "Generated-token rows recover permission-slip, so the distribution mismatch hypothesis was partly real.",
            "All-position rows do not fix tracking and can dilute the generated-token permission signal.",
            "Next Qwen sparse work should be family-specific or transcoder-style, not wider vanilla SAE.",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Gemma Candidate")
    add_bullets(
        slide,
        [
            "Base: google/gemma-2-2b-it; abliterated: IlyaGusev/gemma-2-2b-it-abliterated.",
            "Behavior screen: base harmful clean 1.000, abliterated harmful clean 0.000; both benign helpful 1.000.",
            "Architecture matches: Gemma2, 26 layers, hidden 2304, intermediate 9216.",
            "Public ecosystem: GemmaScope residual/MLP/attention SAEs plus public transcoders.",
        ],
    )

    add_picture_slide(
        prs,
        "Gemma RQ0: Harmful-Specific Activation Drift",
        "gemma_activation_similarity.png",
        "Layer 20 cosine: harmful 0.687 vs benign 0.991.",
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Decision")
    add_bullets(
        slide,
        [
            "Keep Qwen as the distributed-residual case study.",
            "Promote Gemma as the next public SAE/transcoder ecosystem branch.",
            "Run Gemma module/activation patching before GemmaScope feature interpretation.",
            "Proceed to feature naming only after a sparse/transcoder basis beats or explains honest baselines.",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Provenance")
    add_bullets(
        slide,
        [
            "stage3/results/qwen1_5b_residual_sae_generated_full_v0/",
            "stage0/results/candidate_screens_gemma2_2b_abliterated_20260522_clean/",
            "stage1/results/gemma2_2b_abliterated_rq0/",
            "docs/VALUE_ACTION_LESSONS_FOR_MODEL_MERGING.md",
            "Pushed commits: 5114edd and 9ebfbd0.",
        ],
        top=1.7,
    )

    prs.save(ROOT / "qwen_gemma_checkpoint.pptx")


if __name__ == "__main__":
    main()
