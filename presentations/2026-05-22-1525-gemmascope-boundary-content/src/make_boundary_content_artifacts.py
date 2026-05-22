#!/usr/bin/env python3
"""Build figures, comparison tables, and PPTX for the boundary/content checkpoint."""

from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
FIGURES = ROOT / "figures"
PPTX_PATH = ROOT / "gemmascope_boundary_content_checkpoint.pptx"


def clean_variant(model: str) -> str:
    return model.rsplit("feature_subset_", 1)[-1]


def load_deterministic(path: Path, source: str) -> pd.DataFrame:
    df = pd.read_csv(path)
    df = df[(df["condition"] == "deterministic") & (df["layer_group"].isin(["all", "mid_late"]))]
    rows = []
    for _, row in df.iterrows():
        variant = clean_variant(str(row["model"]))
        if variant not in {"mix_decode_delta_abs_k1024", "mix_decode_delta_abs_k2048"}:
            continue
        rows.append(
            {
                "source": source,
                "eval_slice": row["eval_slice"],
                "layer_group": row["layer_group"],
                "budget": variant.replace("mix_decode_delta_abs_", ""),
                "harmful_clean": float(row["harmful_ok_rate"]),
                "unsafe": float(row["harmful_unsafe_continuation_rate"]),
                "benign_helpful": float(row["benign_ok_rate"]),
            }
        )
    return pd.DataFrame(rows)


LOW_CONTENT_EDGE_TOKENS = {"a", "s", "so", "that"}


def is_contentish_token(token: str) -> bool:
    stripped = token.strip()
    if stripped in {"", "\\n", "user", "model", "<end_of_turn>", "<start_of_turn>"}:
        return False
    if stripped.lower() in LOW_CONTENT_EDGE_TOKENS:
        return False
    return any(ch.isalnum() for ch in stripped)


def summarize_event_categories(path: Path) -> tuple[pd.DataFrame, Counter[str]]:
    counts: Counter[str] = Counter()
    tokens: Counter[str] = Counter()
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            event = json.loads(line)
            if event.get("metric") != "abs_delta":
                continue
            token = str(event.get("token", ""))
            tokens[token] += 1
            if is_contentish_token(token):
                counts["content-like token"] += 1
            else:
                counts["assistant-boundary/template"] += 1
    total = sum(counts.values())
    rows = [
        {
            "category": category,
            "count": count,
            "share": count / total if total else 0.0,
        }
        for category, count in counts.items()
    ]
    return pd.DataFrame(rows), tokens


def write_csvs() -> tuple[pd.DataFrame, pd.DataFrame, Counter[str]]:
    all_df = load_deterministic(DATA / "all_token_random_seed_metrics.csv", "all-token selected")
    content_df = load_deterministic(DATA / "content_token_random_seed_metrics.csv", "content-token selected")
    comparison = pd.concat([all_df, content_df], ignore_index=True)
    comparison = comparison.sort_values(["eval_slice", "layer_group", "budget", "source"])
    comparison.to_csv(DATA / "boundary_content_comparison.csv", index=False)

    categories, token_counts = summarize_event_categories(DATA / "all_token_feature_events.jsonl")
    categories.to_csv(DATA / "all_token_abs_delta_event_categories.csv", index=False)
    with (DATA / "all_token_abs_delta_top_tokens.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["token", "count"])
        for token, count in token_counts.most_common(12):
            writer.writerow([token.encode("unicode_escape").decode("ascii"), count])
    return comparison, categories, token_counts


def make_plots(comparison: pd.DataFrame, categories: pd.DataFrame) -> None:
    FIGURES.mkdir(exist_ok=True)

    order = []
    for eval_slice in ["4:8", "8:12"]:
        for group in ["all", "mid_late"]:
            for budget in ["k1024", "k2048"]:
                order.append((eval_slice, group, budget))

    labels = [
        f"{sl}\n{'12-20' if group == 'all' else '15-20'}\n{budget}"
        for sl, group, budget in order
    ]
    x = list(range(len(order)))
    width = 0.36
    colors = {"all-token selected": "#276FBF", "content-token selected": "#C44536"}

    fig, ax = plt.subplots(figsize=(11.2, 5.5))
    for offset, source in [(-width / 2, "all-token selected"), (width / 2, "content-token selected")]:
        values = []
        for eval_slice, group, budget in order:
            row = comparison[
                (comparison["source"] == source)
                & (comparison["eval_slice"] == eval_slice)
                & (comparison["layer_group"] == group)
                & (comparison["budget"] == budget)
            ]
            values.append(float(row.iloc[0]["harmful_clean"]))
        ax.bar([i + offset for i in x], values, width=width, color=colors[source], label=source)

    ax.set_ylim(0, 1.08)
    ax.set_ylabel("Harmful clean refusal rate")
    ax.set_title("Content-token feature selection loses most sparse repair")
    ax.set_xticks(x)
    ax.set_xticklabels(labels, fontsize=9)
    ax.grid(axis="y", alpha=0.25)
    ax.legend(loc="upper right")
    for spine in ["top", "right"]:
        ax.spines[spine].set_visible(False)
    fig.tight_layout()
    fig.savefig(FIGURES / "all_vs_content_harmful_clean.png", dpi=220)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(7.8, 4.4))
    categories = categories.sort_values("count", ascending=False)
    bars = ax.bar(categories["category"], categories["share"] * 100, color=["#5B8E7D", "#D17A22"])
    ax.set_ylabel("Share of top abs-delta event rows (%)")
    ax.set_title("Top feature-delta events are mostly assistant-boundary tokens")
    ax.set_ylim(0, 105)
    ax.grid(axis="y", alpha=0.25)
    for bar, value in zip(bars, categories["share"] * 100):
        ax.text(bar.get_x() + bar.get_width() / 2, value + 2, f"{value:.1f}%", ha="center", fontsize=11)
    for spine in ["top", "right"]:
        ax.spines[spine].set_visible(False)
    fig.tight_layout()
    fig.savefig(FIGURES / "event_category_share.png", dpi=220)
    plt.close(fig)


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    title_shape = slide.shapes.title
    if title_shape is None:
        title_shape = slide.shapes.add_textbox(Inches(0.62), Inches(0.35), Inches(12.0), Inches(0.55))
    title_shape.text = title
    title_shape.text_frame.paragraphs[0].font.size = Pt(30)
    title_shape.text_frame.paragraphs[0].font.bold = True
    if subtitle:
        box = slide.shapes.add_textbox(Inches(0.75), Inches(1.45), Inches(11.8), Inches(0.55))
        p = box.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(15)
        p.font.color.rgb = RGBColor(90, 90, 90)


def add_bullets(slide, bullets: list[str], top: float = 1.65, left: float = 0.85, width: float = 11.5) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(4.8))
    frame = box.text_frame
    frame.clear()
    for idx, bullet in enumerate(bullets):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(18)
        p.space_after = Pt(9)


def add_footer(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.6), Inches(7.08), Inches(12.1), Inches(0.24))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(8.5)
    p.font.color.rgb = RGBColor(115, 115, 115)
    p.alignment = PP_ALIGN.RIGHT


def make_pptx(comparison: pd.DataFrame, categories: pd.DataFrame, token_counts: Counter[str]) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    title_layout = prs.slide_layouts[5]

    slide = prs.slides.add_slide(title_layout)
    add_title(slide, "GemmaScope boundary-vs-content checkpoint", "Model-merging SAE project, 2026-05-22")
    add_bullets(
        slide,
        [
            "Question: are the successful sparse SAE patches content semantics, or response-start state?",
            "Short answer: current evidence points to response-boundary refusal-state transfer.",
            "This is a reframe, not a dead end: it gives a sharper causal test for the next step.",
        ],
        top=2.15,
    )
    add_footer(slide, "Source package: presentations/2026-05-22-1525-gemmascope-boundary-content")

    slide = prs.slides.add_slide(title_layout)
    add_title(slide, "Plain-language setup")
    add_bullets(
        slide,
        [
            "Model merging: combine two or more trained model weight sets without retraining on all original data.",
            "Our pair: a normal Gemma-2-2B-it donor and an abliterated Gemma-2-2B-it recipient.",
            "Patch test: temporarily replace selected internal recipient activations with donor-like SAE reconstructions.",
            "SAE feature: a sparse coordinate in a learned dictionary; useful only if it explains or controls behavior.",
        ],
    )
    add_footer(slide, "Terms are simplified for decision-making; exact artifacts are listed on the provenance slide.")

    slide = prs.slides.add_slide(title_layout)
    add_title(slide, "Where this fits in the project")
    add_bullets(
        slide,
        [
            "Known: full donor post-FF patch over layers 12-20 repairs harmful refusal while keeping benign helpfulness.",
            "Known: full decoded GemmaScope MLP-SAE patch over 12-20 also repairs behavior.",
            "Recent: top-delta feature subsets beat random active-feature controls, especially all 12-20 and 15-20.",
            "New: the top-delta events mostly occur at the assistant response boundary, not ordinary prompt words.",
        ],
    )
    add_footer(slide, "This changes the proposed mechanism we should test next.")

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Main quantitative comparison")
    slide.shapes.add_picture(str(FIGURES / "all_vs_content_harmful_clean.png"), Inches(0.65), Inches(1.22), width=Inches(12.05))
    add_footer(slide, "Y-axis: harmful clean refusal rate on heldout prompts; higher is better.")

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Event audit")
    slide.shapes.add_picture(str(FIGURES / "event_category_share.png"), Inches(0.75), Inches(1.35), width=Inches(6.3))
    top_tokens = ", ".join(
        token.encode("unicode_escape").decode("ascii") for token, _ in token_counts.most_common(5)
    )
    add_bullets(
        slide,
        [
            "All-token audit: 2817 / 2880 top abs-delta rows are boundary/template tokens.",
            f"Most frequent top tokens: {top_tokens}.",
            "Content-selected features activate on harmful words, but the causal repair mostly disappears.",
        ],
        top=1.72,
        left=7.25,
        width=5.2,
    )
    add_footer(slide, "Boundary means role/special/newline tokens around the assistant response start.")

    slide = prs.slides.add_slide(title_layout)
    add_title(slide, "Interpretation")
    add_bullets(
        slide,
        [
            "Result: strongest sparse repair is not yet a clean harmful-content semantic feature story.",
            "Hypothesis: the merge/abliteration weakens a donor-like refusal-mode setup at assistant start.",
            "Speculation: prompt harmfulness features may exist, but they are weak causal handles unless response state is set.",
            "Caveat: current filter changes feature selection; it does not yet prove the runtime patch location.",
        ],
    )
    add_footer(slide, "The caveat is why the next experiment is position-restricted patching.")

    slide = prs.slides.add_slide(title_layout)
    add_title(slide, "Next decisive test")
    add_bullets(
        slide,
        [
            "Use the same successful all-token selected features.",
            "Patch only assistant-boundary positions, only content positions, and all positions.",
            "If boundary-only works and content-only fails, the paper claim becomes response-boundary state transfer.",
            "If content-only works, the current audit was misleading and semantic prompt features remain viable.",
        ],
    )
    add_footer(slide, "This is the highest-value next branch before broader sweeps.")

    slide = prs.slides.add_slide(title_layout)
    add_title(slide, "Provenance")
    add_bullets(
        slide,
        [
            "Finding memo: stage3/results/GEMMA2_2B_GEMMASCOPE_MLP_SAE_BOUNDARY_VS_CONTENT_FINDINGS.md",
            "Source metrics: data/all_token_random_seed_metrics.csv and data/content_token_random_seed_metrics.csv",
            "Source events: data/all_token_feature_events.jsonl and data/content_token_feature_events.jsonl",
            "Generated tables: data/boundary_content_comparison.csv and data/all_token_abs_delta_event_categories.csv",
            "Deck source: gemmascope_boundary_content_checkpoint.tex and this PPTX package.",
        ],
        top=1.35,
    )
    add_footer(slide, "All paths above are local to the repository/package.")

    prs.save(PPTX_PATH)


def main() -> int:
    comparison, categories, token_counts = write_csvs()
    make_plots(comparison, categories)
    make_pptx(comparison, categories, token_counts)
    print(f"[csv] {DATA / 'boundary_content_comparison.csv'}")
    print(f"[csv] {DATA / 'all_token_abs_delta_event_categories.csv'}")
    print(f"[fig] {FIGURES / 'all_vs_content_harmful_clean.png'}")
    print(f"[fig] {FIGURES / 'event_category_share.png'}")
    print(f"[pptx] {PPTX_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
