#!/usr/bin/env python3
"""Build figure and native PPTX for the GemmaScope position-budget deck."""

from __future__ import annotations

import csv
from pathlib import Path

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
FIG_DIR = ROOT / "figures"


def load_rows() -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for path in sorted(DATA_DIR.glob("*metrics.csv")):
        with path.open(newline="", encoding="utf-8") as handle:
            for row in csv.DictReader(handle):
                row["deck_source_csv"] = path.name
                rows.append(row)
    return rows


def value(
    rows: list[dict[str, str]], budget: str, eval_slice: str, patch_filter: str
) -> float | None:
    for row in rows:
        if (
            row["budget"] == budget
            and row["eval_slice"] == eval_slice
            and row["patch_token_filter"] == patch_filter
            and row["layer_group"] == "all"
        ):
            return float(row["harmful_clean"])
    return None


def make_plot(rows: list[dict[str, str]]) -> Path:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    plot_path = FIG_DIR / "position_budget_harmful_clean.png"
    conditions = [
        ("k1024 all", "k1024", "all"),
        ("k1024 bnd+gen", "k1024", "assistant_boundary_or_generated"),
        ("k1024 content+gen", "k1024", "contentish_or_generated"),
        ("k2048 all", "k2048", "all"),
        ("k2048 bnd+gen", "k2048", "assistant_boundary_or_generated"),
        ("k2048 content+gen", "k2048", "contentish_or_generated"),
        ("k512 bnd+gen", "k512", "assistant_boundary_or_generated"),
    ]
    x = list(range(len(conditions)))
    width = 0.36
    folds = [("4:8", -width / 2, "#2b6cb0"), ("8:12", width / 2, "#dd6b20")]
    fig, ax = plt.subplots(figsize=(10.5, 4.6))
    for fold, offset, color in folds:
        vals = [value(rows, budget, fold, filt) for _, budget, filt in conditions]
        heights = [0.0 if val is None else val for val in vals]
        bars = ax.bar([idx + offset for idx in x], heights, width, label=fold, color=color)
        for bar, val in zip(bars, vals, strict=True):
            if val is None:
                ax.text(
                    bar.get_x() + bar.get_width() / 2,
                    0.03,
                    "missing",
                    ha="center",
                    va="bottom",
                    rotation=90,
                    fontsize=7,
                    color="#555555",
                )
            else:
                ax.text(
                    bar.get_x() + bar.get_width() / 2,
                    val + 0.025,
                    f"{val:.2f}",
                    ha="center",
                    va="bottom",
                    fontsize=8,
                )
    ax.set_ylabel("Harmful clean refusal rate")
    ax.set_ylim(0, 1.12)
    ax.set_xticks(x)
    ax.set_xticklabels([label for label, _, _ in conditions], rotation=25, ha="right")
    ax.set_title("Boundary/template plus generated-token patching is the stable reduced path")
    ax.legend(title="Heldout fold")
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(plot_path, dpi=180)
    plt.close(fig)
    return plot_path


def add_title(slide, title: str, subtitle: str = "") -> None:
    title_box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.6))
    title_frame = title_box.text_frame
    title_frame.clear()
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.48), Inches(0.9), Inches(12.2), Inches(0.35))
        sub_frame = sub_box.text_frame
        sub_frame.clear()
        p = sub_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(13)


def add_bullets(slide, bullets: list[str], top: float = 1.35) -> None:
    box = slide.shapes.add_textbox(Inches(0.75), Inches(top), Inches(11.9), Inches(5.3))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    for idx, item in enumerate(bullets):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(21)
        p.space_after = Pt(8)


def add_table(slide, rows: list[list[str]], left: float, top: float, width: float, height: float) -> None:
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width), Inches(height))
    table = table_shape.table
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = text
            for para in cell.text_frame.paragraphs:
                para.alignment = PP_ALIGN.CENTER
                for run in para.runs:
                    run.font.size = Pt(10 if r_idx else 11)
                    run.font.bold = r_idx == 0


def make_pptx(plot_path: Path) -> Path:
    pptx_path = ROOT / "gemmascope_position_budget_checkpoint.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    add_title(slide, "GemmaScope Position/Budget Checkpoint", "Model merging SAE project - 2026-05-22")
    add_bullets(
        slide,
        [
            "Question: where does the sparse repair actually act?",
            "Result: not harmful-content words alone, not a static boundary token alone.",
            "Current mechanism: assistant/template state seeded at response start, then maintained through generated-token history.",
        ],
    )

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Plain Definitions")
    add_bullets(
        slide,
        [
            "Model merging combines model weights or deltas; here the merge damaged refusal behavior.",
            "SAE features are sparse directions used to reconstruct MLP activations in Gemma-2-2B.",
            "A patch replaces selected recipient activations with donor-like decoded SAE contributions.",
            "A position mask asks whether the patch must happen on prompt content, assistant boundary/template tokens, or generated tokens.",
        ],
    )

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Key Evidence")
    add_table(
        slide,
        [
            ["Budget", "Fold", "All pos", "Boundary+gen", "Content+gen", "Template+gen"],
            ["k1024", "4:8", "1.00 / 0.00", "0.75 / 0.00", "0.00 / 0.00", "0.75 / 0.00"],
            ["k1024", "8:12", "0.75 / 0.00", "0.75 / 0.00", "0.00 / 0.00", "0.75 / 0.00"],
            ["k2048", "4:8", "1.00 / 0.00", "0.75 / 0.00", "0.25 / 0.25", "0.75 / 0.00"],
            ["k2048", "8:12", "0.75 / 0.00", "0.75 / 0.25", "missing", "0.75 / 0.25"],
            ["k512", "4:8", "-", "0.75 / 0.00", "-", "-"],
            ["k512", "8:12", "-", "0.50 / 0.00", "-", "-"],
        ],
        left=0.55,
        top=1.35,
        width=12.2,
        height=3.3,
    )
    add_bullets(slide, ["Cells are harmful clean refusal / unsafe continuation."], top=5.15)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Plot View")
    slide.shapes.add_picture(str(plot_path), Inches(0.8), Inches(1.25), width=Inches(11.8))

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Interpretation")
    add_bullets(
        slide,
        [
            "The useful prompt-side seed is template/boundary state, not ordinary harmful-content tokens.",
            "Generated-token history matters: the repair behaves like trajectory maintenance during rollout.",
            "k2048 does not rescue content+generated; it adds some repair but also unsafe continuation.",
            "k512 still works partially, so the causal set is not just an overlarge k1024 accident.",
        ],
    )

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Next Step")
    add_bullets(
        slide,
        [
            "Move from position masks to feature-ID causality inside boundary/template plus generated-token patching.",
            "Ablate or isolate top features by layer and by generated-token time slice.",
            "Goal: identify whether a compact feature bundle carries refusal-state setup and maintenance.",
        ],
    )

    prs.save(pptx_path)
    return pptx_path


def main() -> int:
    rows = load_rows()
    plot_path = make_plot(rows)
    pptx_path = make_pptx(plot_path)
    print(plot_path)
    print(pptx_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
