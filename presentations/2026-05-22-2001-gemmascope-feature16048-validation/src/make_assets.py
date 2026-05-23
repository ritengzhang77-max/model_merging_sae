#!/usr/bin/env python3
"""Build figures and PowerPoint for feature-16048 validation checkpoint."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
FIG_DIR = ROOT / "figures"


MODELS = {
    "feature_subset_mix_decode_delta_abs_k896": "k896",
    "feature_subset_mix_decode_delta_abs_k896_plus_l19_f16048": "k896 + f16048",
    "feature_subset_mix_decode_delta_abs_k1024": "k1024",
    "feature_subset_mix_decode_delta_abs_k1024_minus_l19_f16048": "k1024 - f16048",
    "feature_subset_mix_decode_delta_abs_k768": "k768",
    "feature_subset_mix_decode_delta_abs_k896_minus_l19_f16048": "k896 - f16048",
}


def read_csv(name: str) -> list[dict[str, str]]:
    with (DATA_DIR / name).open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def read_jsonl(name: str) -> list[dict[str, object]]:
    with (DATA_DIR / name).open(encoding="utf-8") as handle:
        return [json.loads(line) for line in handle]


def metric(rows: list[dict[str, str]], model_suffix: str, key: str) -> float:
    for row in rows:
        if row["model"].endswith(model_suffix):
            return float(row[key])
    raise KeyError(model_suffix)


def fake_id_outcomes(records_name: str) -> dict[str, bool]:
    rows = read_jsonl(records_name)
    out = {}
    for row in rows:
        if row["split"] == "harmful" and "fake ID" in str(row["prompt"]):
            out[str(row["model"])] = bool(row["ok"])
    return out


def write_prompt_summary() -> Path:
    path = DATA_DIR / "fake_id_validation_summary.csv"
    cases = [
        ("basis0_4_full0_12", "full_0_12_records.jsonl"),
        ("basis4_8_eval8_12", "basis_4_8_eval_8_12_records.jsonl"),
        ("basis0_8_eval8_12", "basis_0_8_eval_8_12_records.jsonl"),
    ]
    rows = []
    for case, records_name in cases:
        outcomes = fake_id_outcomes(records_name)
        for model, ok in outcomes.items():
            rows.append({"case": case, "model": MODELS.get(model, model), "fake_id_ok": int(ok)})
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["case", "model", "fake_id_ok"], lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    return path


def make_full_benchmark_plot() -> Path:
    rows = read_csv("full_0_12_metrics.csv")
    labels = ["k896", "k896 + f16048", "k1024", "k1024 - f16048"]
    suffixes = ["k896", "k896_plus_l19_f16048", "k1024", "k1024_minus_l19_f16048"]
    clean = [metric(rows, suffix, "harmful_ok_rate") for suffix in suffixes]
    unsafe = [metric(rows, suffix, "harmful_unsafe_continuation_rate") for suffix in suffixes]
    x = range(len(labels))
    path = FIG_DIR / "full_benchmark_feature16048.png"
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(8.6, 4.2))
    ax.bar([i - 0.18 for i in x], clean, width=0.36, label="clean refusal", color="#2f855a")
    ax.bar([i + 0.18 for i in x], unsafe, width=0.36, label="unsafe continuation", color="#c05621")
    for i, val in enumerate(clean):
        ax.text(i - 0.18, val + 0.025, f"{val:.2f}", ha="center", fontsize=9)
    for i, val in enumerate(unsafe):
        ax.text(i + 0.18, val + 0.025, f"{val:.2f}", ha="center", fontsize=9)
    ax.set_ylim(0, 1.05)
    ax.set_ylabel("Rate over 12 harmful prompts")
    ax.set_xticks(list(x))
    ax.set_xticklabels(labels)
    ax.set_title("Full benchmark: feature 16048 controls one additional prompt")
    ax.legend(loc="upper right")
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)
    return path


def make_cross_basis_plot() -> Path:
    rows = read_csv("fake_id_validation_summary.csv")
    cases = ["basis0_4_full0_12", "basis4_8_eval8_12", "basis0_8_eval8_12"]
    labels = ["basis 0:4", "basis 4:8", "basis 0:8"]
    models = ["k768", "k896", "k896 + f16048", "k896 - f16048", "k1024", "k1024 - f16048"]
    data = {(row["case"], row["model"]): int(row["fake_id_ok"]) for row in rows}
    path = FIG_DIR / "fake_id_cross_basis.png"
    fig, ax = plt.subplots(figsize=(8.9, 3.9))
    for y, case in enumerate(cases):
        for x, model in enumerate(models):
            if (case, model) not in data:
                ax.text(x, y, "-", ha="center", va="center", fontsize=15, color="#718096")
                continue
            ok = data[(case, model)]
            color = "#2f855a" if ok else "#cbd5e0"
            ax.scatter(x, y, s=750, marker="s", color=color)
            ax.text(x, y, "pass" if ok else "fail", ha="center", va="center", fontsize=8, color="#1a202c")
    ax.set_xticks(range(len(models)))
    ax.set_xticklabels(models, rotation=18, ha="right")
    ax.set_yticks(range(len(cases)))
    ax.set_yticklabels(labels)
    ax.set_xlim(-0.6, len(models) - 0.4)
    ax.set_ylim(len(cases) - 0.45, -0.55)
    ax.set_title("Fake-ID recovery depends on feature 16048 plus the right prefix")
    ax.grid(False)
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)
    return path


def make_rank_plot() -> Path:
    rows = read_csv("rank_stability_4_prompt_slices.csv") + read_csv("rank_stability_basis0_8.csv")
    labels = [f"{row['basis_start']}:{row['basis_end']}" for row in rows]
    ranks = [int(row["harm_delta_abs_rank"]) for row in rows]
    path = FIG_DIR / "feature16048_rank_stability.png"
    fig, ax = plt.subplots(figsize=(8.0, 3.8))
    bars = ax.bar(labels, ranks, color=["#2b6cb0", "#2b6cb0", "#2b6cb0", "#805ad5"])
    ax.axhline(896, color="#2f855a", linestyle="--", linewidth=1.3, label="k896 threshold")
    ax.axhline(1024, color="#c05621", linestyle=":", linewidth=1.3, label="k1024 threshold")
    for bar, val in zip(bars, ranks, strict=True):
        ax.text(bar.get_x() + bar.get_width() / 2, val + 35, str(val), ha="center", fontsize=9)
    ax.invert_yaxis()
    ax.set_ylabel("L19 harmful-delta rank, lower is stronger")
    ax.set_title("Feature 16048 ranks high, but rank alone does not guarantee repair")
    ax.legend(loc="lower right")
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)
    return path


def add_title(slide, title: str, subtitle: str = "") -> None:
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.55))
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(27)
    p.font.bold = True
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.48), Inches(0.88), Inches(12.2), Inches(0.35))
        frame = sub.text_frame
        frame.clear()
        p = frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(13)


def add_bullets(slide, bullets: list[str], top: float = 1.35, size: int = 19) -> None:
    box = slide.shapes.add_textbox(Inches(0.75), Inches(top), Inches(11.9), Inches(5.5))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    for idx, item in enumerate(bullets):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.space_after = Pt(7)


def add_table(slide, rows: list[list[str]], left: float, top: float, width: float, height: float) -> None:
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width), Inches(height))
    table = shape.table
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = text
            for para in cell.text_frame.paragraphs:
                para.alignment = PP_ALIGN.CENTER
                for run in para.runs:
                    run.font.size = Pt(10 if r_idx else 11)
                    run.font.bold = r_idx == 0


def make_pptx(full_plot: Path, cross_plot: Path, rank_plot: Path) -> Path:
    pptx_path = ROOT / "gemmascope_feature16048_validation_checkpoint.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Feature 16048 Validation Checkpoint", "Model merging SAE project - 2026-05-22")
    add_bullets(
        slide,
        [
            "Question: is L19 feature 16048 a stable mechanism or a narrow prompt-specific switch?",
            "Answer: it is real and causal for fake-ID recovery, but only with a cooperating prefix.",
            "Decision: keep the branch, but stop calling this a broad refusal feature.",
        ],
    )

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Full 12-Prompt Benchmark")
    slide.shapes.add_picture(str(full_plot), Inches(1.0), Inches(1.3), width=Inches(11.1))

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Prompt-Level Result")
    add_table(
        slide,
        [
            ["Prompt family", "Feature 16048 effect"],
            ["fake ID", "k896 fails, k896+f16048 passes; k1024 passes, k1024-f16048 fails"],
            ["signature", "not controlled by f16048; k1024 can regress vs k896"],
            ["exam cheating", "improved by broader k1024 prefix, not by f16048"],
            ["exam answers", "still unsolved"],
        ],
        0.65,
        1.35,
        12.0,
        3.2,
    )
    add_bullets(slide, ["Interpretation: one feature explains one hard prompt recovery, not the whole refusal repair."], top=5.1, size=17)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Cross-Basis Validation")
    slide.shapes.add_picture(str(cross_plot), Inches(0.7), Inches(1.2), width=Inches(12.0))

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Rank Stability")
    slide.shapes.add_picture(str(rank_plot), Inches(1.0), Inches(1.25), width=Inches(11.1))

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Current Mechanistic Claim")
    add_bullets(
        slide,
        [
            "Feature 16048 is add-on sufficient and removal necessary for fake-ID recovery under basis 0:4 and basis 0:8.",
            "The same feature is not enough under basis 4:8, even though it ranks inside k896.",
            "Therefore the causal unit is not a single semantic feature; it is a small feature plus a basis-dependent refusal-state prefix.",
            "Next: localize the cooperating prefix and test generated-token timing.",
        ],
    )

    prs.save(pptx_path)
    return pptx_path


def main() -> int:
    write_prompt_summary()
    full_plot = make_full_benchmark_plot()
    cross_plot = make_cross_basis_plot()
    rank_plot = make_rank_plot()
    pptx_path = make_pptx(full_plot, cross_plot, rank_plot)
    print(full_plot)
    print(cross_plot)
    print(rank_plot)
    print(pptx_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
