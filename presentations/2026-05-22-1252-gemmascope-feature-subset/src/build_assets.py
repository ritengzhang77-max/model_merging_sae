#!/usr/bin/env python3
"""Build figures and a native PPTX for the GemmaScope feature-subset checkpoint."""

from __future__ import annotations

import csv
from pathlib import Path

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
FIGURES = ROOT / "figures"


def load_metrics(name: str) -> dict[str, dict[str, float | str]]:
    path = DATA / name
    with path.open(newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    return {str(row["model"]): row for row in rows}


def f(row: dict[str, float | str], key: str) -> float:
    return float(row[key])


def build_plot() -> None:
    FIGURES.mkdir(parents=True, exist_ok=True)
    fold_a = load_metrics("heldout_4_8_metrics.csv")
    fold_b = load_metrics("heldout_8_12_metrics.csv")
    panels = [
        (
            "Eval prompts 4-7",
            fold_a,
            [
                ("full SAE", "feature_subset_full_decode"),
                ("all delta", "feature_subset_delta_add_all"),
                ("top k1024", "feature_subset_mix_decode_delta_abs_k1024"),
                ("rand k2048", "feature_subset_mix_decode_random_active_k2048"),
            ],
        ),
        (
            "Eval prompts 8-11",
            fold_b,
            [
                ("full SAE", "feature_subset_full_decode"),
                ("all delta", "feature_subset_delta_add_all"),
                ("top k1024", "feature_subset_mix_decode_delta_abs_k1024"),
                ("rand k1024", "feature_subset_mix_decode_random_active_k1024"),
                ("rand k2048", "feature_subset_mix_decode_random_active_k2048"),
            ],
        ),
    ]
    fig, axes = plt.subplots(1, 2, figsize=(10.5, 3.8), sharey=True)
    colors = ["#4c78a8", "#72b7b2", "#54a24b", "#b279a2", "#e45756"]
    for ax, (title, rows, variants) in zip(axes, panels):
        vals = [f(rows[key], "harmful_ok_rate") for _label, key in variants]
        labels = [label for label, _key in variants]
        ax.bar(labels, vals, color=colors[: len(vals)], width=0.72)
        ax.set_title(title, fontsize=12)
        ax.set_ylim(0, 1.05)
        ax.set_ylabel("harmful clean refusal rate")
        ax.grid(axis="y", alpha=0.25)
        ax.tick_params(axis="x", labelrotation=22)
        for idx, val in enumerate(vals):
            ax.text(idx, val + 0.035, f"{val:.2f}", ha="center", va="bottom", fontsize=9)
    fig.suptitle("GemmaScope MLP SAE feature subsets: heldout refusal repair", fontsize=14)
    fig.tight_layout(rect=(0, 0, 1, 0.92))
    fig.savefig(FIGURES / "heldout_feature_subset_rates.png", dpi=220)
    plt.close(fig)


def add_title(slide, title: str, subtitle: str = ""):
    slide.shapes.title.text = title
    if subtitle:
        box = slide.shapes.add_textbox(Inches(0.9), Inches(1.55), Inches(11.5), Inches(1.0))
        tf = box.text_frame
        tf.text = subtitle
        tf.paragraphs[0].font.size = Pt(23)


def add_bullets(slide, items: list[str], top=1.55):
    box = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11.5), Inches(4.8))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(22)


def build_pptx() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    add_title(
        slide,
        "Model Merging SAE Checkpoint",
        "GemmaScope MLP feature subsets for refusal repair, 2026-05-22",
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Plain Terms")
    add_bullets(
        slide,
        [
            "Model merging here means moving behavior from a donor/base model into an abliterated recipient by patching internal activations.",
            "An SAE is a sparse autoencoder: it rewrites an activation as many mostly-zero feature coordinates, then decodes them back.",
            "The key question is whether specific SAE coordinates explain the repair, or whether only the full dense reconstruction works.",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Current Claim Test")
    add_bullets(
        slide,
        [
            "Target models: google/gemma-2-2b-it donor and IlyaGusev/gemma-2-2b-it-abliterated recipient.",
            "Patch site: normalized post-feedforward MLP update over layers 12-20.",
            "Selector: top SAE coordinates by harmful donor-recipient activation delta.",
            "Control: same-size random active SAE coordinates.",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Heldout Results")
    slide.shapes.add_picture(str(FIGURES / "heldout_feature_subset_rates.png"), Inches(0.75), Inches(1.35), width=Inches(11.9))

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Interpretation")
    add_bullets(
        slide,
        [
            "Positive signal: top-delta SAE coordinates beat matched random controls on heldout prompts.",
            "Not complete: k1024 per layer still fails one heldout harmful family on prompts 8-11.",
            "Best current wording: feature-coordinate causal evidence, not yet a human-interpretable mechanism.",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Next Work")
    add_bullets(
        slide,
        [
            "Localize the feature repair by layer group instead of patching all layers 12-20.",
            "Repeat random controls across seeds at k1024 and k2048.",
            "Export feature IDs, activation deltas, and top activating snippets for interpretability audit.",
            "Stratify failures by prompt family to learn why the exam-answer case is still weak.",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Local Evidence")
    add_bullets(
        slide,
        [
            "stage3/results/GEMMA2_2B_GEMMASCOPE_MLP_SAE_FEATURE_SUBSET_FINDINGS.md",
            "stage3/results/gemma2_2b_gemmascope_mlp_sae_feature_subsets_12_20_heldout_k_sweep_v0/",
            "stage3/results/gemma2_2b_gemmascope_mlp_sae_feature_subsets_12_20_heldout_fold2_v0/",
            "stage3/scripts/run_gemma2_2b_gemmascope_mlp_sae_feature_subsets.py",
        ],
        top=1.45,
    )

    prs.save(ROOT / "gemmascope_feature_subset_checkpoint.pptx")


def main() -> int:
    build_plot()
    build_pptx()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
