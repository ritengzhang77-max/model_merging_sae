#!/usr/bin/env python3
"""Build figures and PPTX for the position-restricted GemmaScope checkpoint."""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
FIGURES = ROOT / "figures"
PPTX = ROOT / "gemmascope_position_restricted_checkpoint.pptx"


def load() -> pd.DataFrame:
    return pd.read_csv(DATA / "position_restricted_metrics.csv")


def make_plots(df: pd.DataFrame) -> None:
    FIGURES.mkdir(exist_ok=True)
    focus = df[(df["budget"] == "k1024") & (df["layer_group"] == "all")]
    order = [
        "all",
        "assistant_boundary",
        "contentish",
        "generated",
        "prompt_all",
        "prompt_or_last",
        "assistant_boundary_or_generated",
        "contentish_or_generated",
        "prompt_template_or_generated",
    ]
    labels = [
        "all",
        "boundary",
        "content",
        "generated",
        "prompt all",
        "prompt+last",
        "boundary+gen",
        "content+gen",
        "template+gen",
    ]
    colors = {
        "all": "#276FBF",
        "assistant_boundary": "#D17A22",
        "contentish": "#C44536",
        "generated": "#777777",
        "prompt_all": "#8E6C8A",
        "prompt_or_last": "#8064A2",
        "assistant_boundary_or_generated": "#5B8E7D",
        "contentish_or_generated": "#B23A48",
        "prompt_template_or_generated": "#3E885B",
    }
    for eval_slice in ["4:8", "8:12"]:
        sub = focus[focus["eval_slice"] == eval_slice]
        vals = []
        shown_labels = []
        shown_colors = []
        for patch, label in zip(order, labels):
            row = sub[sub["patch_token_filter"] == patch]
            if row.empty:
                continue
            vals.append(float(row.iloc[0]["harmful_clean"]))
            shown_labels.append(label)
            shown_colors.append(colors[patch])
        fig, ax = plt.subplots(figsize=(10.6, 4.8))
        ax.bar(range(len(vals)), vals, color=shown_colors)
        ax.set_ylim(0, 1.05)
        ax.set_ylabel("Harmful clean refusal rate")
        ax.set_title(f"All layers 12-20, heldout {eval_slice}, k1024")
        ax.set_xticks(range(len(vals)))
        ax.set_xticklabels(shown_labels, rotation=25, ha="right")
        ax.grid(axis="y", alpha=0.25)
        for i, val in enumerate(vals):
            ax.text(i, val + 0.03, f"{val:.2f}", ha="center", fontsize=10)
        for spine in ["top", "right"]:
            ax.spines[spine].set_visible(False)
        fig.tight_layout()
        fig.savefig(FIGURES / f"all_12_20_{eval_slice.replace(':', '_')}_position_modes.png", dpi=220)
        plt.close(fig)


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    title_shape = slide.shapes.add_textbox(Inches(0.62), Inches(0.35), Inches(12.0), Inches(0.6))
    p = title_shape.text_frame.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(30)
    if subtitle:
        box = slide.shapes.add_textbox(Inches(0.75), Inches(1.13), Inches(11.7), Inches(0.4))
        p = box.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(14)
        p.font.color.rgb = RGBColor(90, 90, 90)


def add_bullets(slide, bullets: list[str], top: float = 1.55) -> None:
    box = slide.shapes.add_textbox(Inches(0.82), Inches(top), Inches(11.8), Inches(5.3))
    frame = box.text_frame
    frame.clear()
    for idx, bullet in enumerate(bullets):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = bullet
        p.font.size = Pt(18)
        p.space_after = Pt(9)


def add_footer(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12.1), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(8.5)
    p.font.color.rgb = RGBColor(115, 115, 115)


def make_pptx() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    add_title(slide, "GemmaScope position-restricted checkpoint", "Model-merging SAE project, 2026-05-22")
    add_bullets(
        slide,
        [
            "Question: where must the same selected SAE features be patched to restore refusal?",
            "Short answer: not content tokens, and not a static assistant boundary alone.",
            "Best reduced mechanism: assistant boundary seed plus generated-token trajectory maintenance.",
        ],
        top=1.75,
    )
    add_footer(slide, "Package: presentations/2026-05-22-1644-gemmascope-position-restricted")

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Why this changes the claim")
    add_bullets(
        slide,
        [
            "Earlier audit: top feature deltas were concentrated on assistant-boundary/template tokens.",
            "New causal test: boundary-only patching fails, so the audit location alone is not sufficient.",
            "Generated-token history alone also fails, so rollout maintenance needs a prompt-side seed.",
            "Content-token conditions remain weak, including content plus generated history.",
        ],
    )
    add_footer(slide, "Same all-token top-delta k1024 features are used across modes.")

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Heldout 4:8")
    slide.shapes.add_picture(str(FIGURES / "all_12_20_4_8_position_modes.png"), Inches(0.75), Inches(1.35), width=Inches(11.8))
    add_footer(slide, "Y-axis: harmful clean refusal rate; benign helpfulness is 1.000 for all rows.")

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Heldout 8:12")
    slide.shapes.add_picture(str(FIGURES / "all_12_20_8_12_position_modes.png"), Inches(0.75), Inches(1.35), width=Inches(11.8))
    add_footer(slide, "Boundary+generated and template+generated match all-position repair on this fold.")

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Interpretation")
    add_bullets(
        slide,
        [
            "Mechanism candidate: an autoregressive refusal-state trajectory.",
            "Prompt-side seed: assistant/template boundary state, not ordinary harmful-content tokens.",
            "Rollout maintenance: selected donor-like SAE coordinates must continue on generated-token history.",
            "Paper framing: model merging can lose a response-state maintenance process, not only a concept feature.",
        ],
    )
    add_footer(slide, "This is stronger than the previous boundary audit because it is causal.")

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Next experiment")
    add_bullets(
        slide,
        [
            "Repeat the same reduced modes at k2048 to check budget sensitivity.",
            "Audit exact top features that survive boundary+generated repair.",
            "Patch generated-token history by time step to find whether refusal needs early or sustained maintenance.",
            "Then move from position causality to feature-ID causality.",
        ],
    )
    add_footer(slide, "Local summary: data/GEMMA2_2B_GEMMASCOPE_MLP_SAE_POSITION_RESTRICTED_SUMMARY.md")

    prs.save(PPTX)


def main() -> int:
    df = load()
    make_plots(df)
    make_pptx()
    print(f"[figures] {FIGURES}")
    print(f"[pptx] {PPTX}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
