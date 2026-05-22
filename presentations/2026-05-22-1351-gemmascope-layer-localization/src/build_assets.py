#!/usr/bin/env python3
"""Build figures and native PPTX for the GemmaScope layer-localization checkpoint."""

from __future__ import annotations

import csv
from pathlib import Path

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
FIGURES = ROOT / "figures"


def load_metrics(filename: str) -> list[dict[str, str]]:
    with (DATA / filename).open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def find_rate(rows: list[dict[str, str]], group: str, suffix: str) -> float:
    for row in rows:
        if row.get("layer_group") == group and row["model"].endswith(suffix):
            return float(row["harmful_ok_rate"])
    raise KeyError((group, suffix))


def build_single_band_plot() -> None:
    FIGURES.mkdir(parents=True, exist_ok=True)
    fold_a = load_metrics("single_4_8_metrics.csv")
    fold_b = load_metrics("single_8_12_metrics.csv")
    groups = ["all", "early", "mid", "late"]
    labels = ["all\n12-20", "early\n12-14", "mid\n15-17", "late\n18-20"]
    suffix = "full_decode"
    vals_a = [find_rate(fold_a, g, suffix) for g in groups]
    vals_b = [find_rate(fold_b, g, suffix) for g in groups]
    x = range(len(groups))
    width = 0.36
    fig, ax = plt.subplots(figsize=(8.8, 4.1))
    ax.bar([i - width / 2 for i in x], vals_a, width=width, label="eval 4-7", color="#4c78a8")
    ax.bar([i + width / 2 for i in x], vals_b, width=width, label="eval 8-11", color="#f58518")
    ax.set_xticks(list(x))
    ax.set_xticklabels(labels)
    ax.set_ylim(0, 1.05)
    ax.set_ylabel("harmful clean refusal rate")
    ax.set_title("Single 3-layer bands are not sufficient")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False)
    for i, val in enumerate(vals_a):
        ax.text(i - width / 2, val + 0.035, f"{val:.2f}", ha="center", fontsize=9)
    for i, val in enumerate(vals_b):
        ax.text(i + width / 2, val + 0.035, f"{val:.2f}", ha="center", fontsize=9)
    fig.tight_layout()
    fig.savefig(FIGURES / "single_band_full_decode.png", dpi=220)
    plt.close(fig)


def build_pair_plot() -> None:
    fold_a = load_metrics("pairs_4_8_metrics.csv")
    fold_b = load_metrics("pairs_8_12_metrics.csv")
    groups = ["early_mid", "mid_late", "early_late"]
    labels = ["12-17", "15-20", "12-14 +\n18-20"]
    suffix = "mix_decode_delta_abs_k2048"
    vals_a = [find_rate(fold_a, g, suffix) for g in groups]
    vals_b = [find_rate(fold_b, g, suffix) for g in groups]
    x = range(len(groups))
    width = 0.36
    fig, ax = plt.subplots(figsize=(8.8, 4.1))
    ax.bar([i - width / 2 for i in x], vals_a, width=width, label="eval 4-7", color="#54a24b")
    ax.bar([i + width / 2 for i in x], vals_b, width=width, label="eval 8-11", color="#e45756")
    ax.set_xticks(list(x))
    ax.set_xticklabels(labels)
    ax.set_ylim(0, 1.05)
    ax.set_ylabel("harmful clean refusal rate")
    ax.set_title("Late-containing pairs recover much more behavior")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False)
    for i, val in enumerate(vals_a):
        ax.text(i - width / 2, val + 0.035, f"{val:.2f}", ha="center", fontsize=9)
    for i, val in enumerate(vals_b):
        ax.text(i + width / 2, val + 0.035, f"{val:.2f}", ha="center", fontsize=9)
    fig.tight_layout()
    fig.savefig(FIGURES / "pair_band_top_delta_k2048.png", dpi=220)
    plt.close(fig)


def add_title(slide, title: str, subtitle: str = ""):
    slide.shapes.title.text = title
    if subtitle:
        box = slide.shapes.add_textbox(Inches(0.85), Inches(1.5), Inches(11.7), Inches(0.9))
        tf = box.text_frame
        tf.text = subtitle
        tf.paragraphs[0].font.size = Pt(22)


def add_bullets(slide, items: list[str], top=1.45):
    box = slide.shapes.add_textbox(Inches(0.85), Inches(top), Inches(11.8), Inches(5.2))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(21)
        p.level = 0


def build_pptx() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    add_title(
        slide,
        "GemmaScope SAE Layer Localization",
        "Where the model-merging refusal repair lives, 2026-05-22",
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Question")
    add_bullets(
        slide,
        [
            "Earlier result: selected GemmaScope MLP-SAE features over layers 12-20 can repair refusal better than random active features.",
            "New question: is the repair concentrated in one small layer band, or distributed across multiple bands?",
            "Operation: patch decoded SAE features into the abliterated model at normalized post-feedforward MLP outputs.",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Single Bands")
    slide.shapes.add_picture(str(FIGURES / "single_band_full_decode.png"), Inches(1.0), Inches(1.25), width=Inches(11.3))

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Paired Bands")
    slide.shapes.add_picture(str(FIGURES / "pair_band_top_delta_k2048.png"), Inches(1.0), Inches(1.25), width=Inches(11.3))

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Interpretation")
    add_bullets(
        slide,
        [
            "No 3-layer band is independently sufficient; late layers 18-20 are strongest but still partial.",
            "The weak 12-17 result means late layers are necessary.",
            "Late-containing six-layer pairs recover much more behavior, but the best pair depends on prompt slice.",
            "Current wording: late-dependent multi-layer composition, not a single compact layer module.",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Next")
    add_bullets(
        slide,
        [
            "Repeat random controls across seeds for all 12-20, 15-20, and 12-14+18-20.",
            "Narrow late-containing groups: 15-18, 16-20, 17-20, 15-17+20, and 12-14+18-19.",
            "Export feature IDs, activation deltas, and top activating snippets before assigning semantic labels.",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "Local Evidence")
    add_bullets(
        slide,
        [
            "stage3/results/GEMMA2_2B_GEMMASCOPE_MLP_SAE_LAYER_GROUP_FINDINGS.md",
            "stage3/scripts/run_gemma2_2b_gemmascope_mlp_sae_layer_groups.py",
            "stage3/results/gemma2_2b_gemmascope_mlp_sae_layer_groups_eval_4_8_v0/",
            "stage3/results/gemma2_2b_gemmascope_mlp_sae_layer_groups_pairs_eval_8_12_v0/",
        ],
        top=1.35,
    )

    prs.save(ROOT / "gemmascope_layer_localization_checkpoint.pptx")


def main() -> int:
    build_single_band_plot()
    build_pair_plot()
    build_pptx()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
