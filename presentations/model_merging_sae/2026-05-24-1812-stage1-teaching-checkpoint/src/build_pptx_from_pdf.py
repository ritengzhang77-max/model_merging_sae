#!/usr/bin/env python3
"""Create a PPTX visual mirror from the compiled Beamer PDF."""

from __future__ import annotations

import re
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


THIS = Path(__file__).resolve()
DECK_DIR = THIS.parents[1]
OUT = DECK_DIR / "outputs"
PDF = OUT / "model_merging_sae_stage1_teaching_checkpoint.pdf"
PPTX = OUT / "model_merging_sae_stage1_teaching_checkpoint.pptx"


def page_number(path: Path) -> int:
    match = re.search(r"-(\d+)\.png$", path.name)
    if not match:
        return 0
    return int(match.group(1))


def main() -> None:
    if not PDF.exists():
        raise SystemExit(f"Missing PDF: {PDF}")

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    with tempfile.TemporaryDirectory(prefix="model_merging_sae_pptx_") as tmp:
        prefix = Path(tmp) / "slide"
        subprocess.run(
            ["pdftoppm", "-png", "-r", "144", str(PDF), str(prefix)],
            check=True,
        )
        images = sorted(Path(tmp).glob("slide-*.png"), key=page_number)
        if not images:
            raise SystemExit("pdftoppm produced no slide images")

        for image in images:
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(
                str(image),
                0,
                0,
                width=prs.slide_width,
                height=prs.slide_height,
            )

    prs.save(PPTX)
    print(PPTX)


if __name__ == "__main__":
    main()
